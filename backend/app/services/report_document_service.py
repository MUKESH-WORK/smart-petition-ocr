"""
Backend Document & Report Generation Service
Supports generating authoritative Government of Tamil Nadu audit reports in PDF and DOCX formats:
1. Officer Performance & Processed Petition Counts / Status
2. System & Officer Audit Log History
3. Particular Petition Full Form Details Dossier
4. Total Petitions Received & Intake Analysis

Typography Requirements:
- Latin/English: Bookman Old Style
- Tamil Script: Marudham / Nirmala UI / Mukta Malar
"""

import io
import os
import json
from datetime import datetime, timezone
from typing import Dict, Any, List, Optional, Tuple

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# Theme Colors
HEX_NAVY = "102C57"
HEX_GOLD = "DAC0A3"
HEX_SLATE = "475569"
HEX_LIGHT_BG = "F8FAFC"
HEX_LABEL_BG = "F1F5F9"
HEX_BORDER = "CBD5E1"
HEX_GREEN = "166534"

RL_NAVY = colors.HexColor("#102C57")
RL_GOLD = colors.HexColor("#DAC0A3")
RL_SLATE = colors.HexColor("#475569")
RL_LIGHT_BG = colors.HexColor("#F8FAFC")
RL_LABEL_BG = colors.HexColor("#F1F5F9")
RL_BORDER = colors.HexColor("#CBD5E1")
RL_GREEN = colors.HexColor("#166534")


# ==============================================================================
# FONT REGISTRATION FOR REPORTLAB (PDF)
# ==============================================================================

PDF_REGULAR_FONT = "Helvetica"
PDF_BOLD_FONT = "Helvetica-Bold"
PDF_TAMIL_FONT = "Helvetica"
PDF_TAMIL_BOLD = "Helvetica-Bold"

def _init_pdf_fonts():
    global PDF_REGULAR_FONT, PDF_BOLD_FONT, PDF_TAMIL_FONT, PDF_TAMIL_BOLD
    
    # 1. Look for Windows fonts with Tamil & Bookman support
    candidates = [
        ("NirmalaUI", "C:/Windows/Fonts/Nirmala.ttf"),
        ("NirmalaUI-Bold", "C:/Windows/Fonts/NirmalaB.ttf"),
        ("BookmanOldStyle", "C:/Windows/Fonts/BOOKOS.TTF"),
        ("BookmanOldStyle-Bold", "C:/Windows/Fonts/BOOKOSB.TTF"),
        ("Latha", "C:/Windows/Fonts/Latha.ttf"),
        ("Latha-Bold", "C:/Windows/Fonts/Lathab.ttf"),
    ]

    registered = {}
    for name, path in candidates:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont(name, path))
                registered[name] = True
            except Exception as e:
                pass

    if "NirmalaUI" in registered:
        PDF_TAMIL_FONT = "NirmalaUI"
        PDF_TAMIL_BOLD = "NirmalaUI-Bold" if "NirmalaUI-Bold" in registered else "NirmalaUI"
        # Nirmala UI contains both Latin (English) and Tamil Unicode
        PDF_REGULAR_FONT = "NirmalaUI"
        PDF_BOLD_FONT = "NirmalaUI-Bold" if "NirmalaUI-Bold" in registered else "NirmalaUI"
    elif "Latha" in registered:
        PDF_TAMIL_FONT = "Latha"
        PDF_TAMIL_BOLD = "Latha-Bold" if "Latha-Bold" in registered else "Latha"

    if "BookmanOldStyle" in registered and PDF_REGULAR_FONT == "Helvetica":
        PDF_REGULAR_FONT = "BookmanOldStyle"
        PDF_BOLD_FONT = "BookmanOldStyle-Bold" if "BookmanOldStyle-Bold" in registered else "BookmanOldStyle"

_init_pdf_fonts()


def _get_pdf_styles():
    base = getSampleStyleSheet()
    
    style_normal = ParagraphStyle(
        'GovtNormal',
        parent=base['Normal'],
        fontName=PDF_REGULAR_FONT,
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#0F172A")
    )
    style_bold = ParagraphStyle(
        'GovtBold',
        parent=style_normal,
        fontName=PDF_BOLD_FONT,
        fontSize=8,
        leading=11
    )
    style_tamil = ParagraphStyle(
        'GovtTamil',
        parent=style_normal,
        fontName=PDF_TAMIL_FONT,
        fontSize=8,
        leading=11.5,
        textColor=colors.HexColor("#0F172A")
    )
    style_tamil_bold = ParagraphStyle(
        'GovtTamilBold',
        parent=style_tamil,
        fontName=PDF_TAMIL_BOLD,
        fontSize=8.5,
        leading=12
    )
    style_heading = ParagraphStyle(
        'GovtHeading',
        parent=style_normal,
        fontName=PDF_BOLD_FONT,
        fontSize=9.5,
        leading=13,
        textColor=RL_NAVY
    )

    return {
        "normal": style_normal,
        "bold": style_bold,
        "tamil": style_tamil,
        "tamil_bold": style_tamil_bold,
        "heading": style_heading
    }


# ==============================================================================
# DOCX STYLING HELPERS
# ==============================================================================

def _set_cell_bg(cell, hex_color: str):
    """Sets cell background color in docx."""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def _set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    """Sets cell padding in docx."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def _set_table_borders(table):
    """Sets subtle light borders on docx tables."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def _format_run_fonts(run, size_pt=9, bold=False, color_rgb=(15, 23, 42)):
    """Applies Bookman Old Style (Latin) and Marudham (Tamil) to run."""
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.color.rgb = RGBColor(*color_rgb)
    run.font.name = "Bookman Old Style"
    
    # Set complex script / EastAsia fonts to Marudham / Nirmala UI
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(
        f'<w:rFonts {nsdecls("w")} '
        f'w:ascii="Bookman Old Style" '
        f'w:hAnsi="Bookman Old Style" '
        f'w:cs="Marudham" '
        f'w:eastAsia="Marudham"/>'
    )
    rPr.append(rFonts)


def _add_docx_header(doc: Document, title: str, subtitle: str = "REVENUE & DISASTER MANAGEMENT DEPARTMENT • ERODE DISTRICT COLLECTORATE"):
    """Adds official Tamil Nadu government header and title block."""
    for s in doc.sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.65)
        s.right_margin = Inches(0.65)
        s.different_first_page_header_footer = False

    # Document Header Title
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.paragraph_format.space_after = Pt(2)

    r1 = h.add_run("GOVERNMENT OF TAMIL NADU\n")
    _format_run_fonts(r1, size_pt=13, bold=True, color_rgb=(16, 44, 87))

    r2 = h.add_run(f"{subtitle}\n")
    _format_run_fonts(r2, size_pt=8.5, bold=True, color_rgb=(71, 85, 105))

    r3 = h.add_run(title)
    _format_run_fonts(r3, size_pt=10.5, bold=True, color_rgb=(16, 44, 87))

    # Divider bar
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(10)
    r_div = p_div.add_run("―" * 58)
    _format_run_fonts(r_div, size_pt=8, bold=True, color_rgb=(218, 192, 163))


def _add_docx_footer_and_seal(doc: Document):
    """Adds official seal, DRO signature block and certified notice."""
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(14)
    p_space.paragraph_format.space_after = Pt(4)

    t_sign = doc.add_table(rows=1, cols=2)
    t_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(t_sign)

    # Left cell: Seal
    cell_seal = t_sign.rows[0].cells[0]
    _set_cell_bg(cell_seal, HEX_LIGHT_BG)
    _set_cell_margins(cell_seal, top=140, bottom=140, left=160, right=160)
    p_seal = cell_seal.paragraphs[0]
    p_seal.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s1 = p_seal.add_run("【 OFFICIAL GOVERNMENT SEAL 】\n")
    _format_run_fonts(r_s1, size_pt=8.5, bold=True, color_rgb=(71, 85, 105))
    r_s2 = p_seal.add_run("DISTRICT COLLECTORATE • ERODE\nPUBLIC GRIEVANCE CELL")
    _format_run_fonts(r_s2, size_pt=7.5, color_rgb=(100, 116, 139))

    # Right cell: Signature
    cell_sign = t_sign.rows[0].cells[1]
    _set_cell_bg(cell_sign, HEX_LIGHT_BG)
    _set_cell_margins(cell_sign, top=140, bottom=140, left=160, right=160)
    p_sign = cell_sign.paragraphs[0]
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_line = p_sign.add_run("________________________________\n")
    _format_run_fonts(r_line, size_pt=8, color_rgb=(148, 163, 184))
    r_sig1 = p_sign.add_run("District Administrator / DRO\n")
    _format_run_fonts(r_sig1, size_pt=9, bold=True, color_rgb=(16, 44, 87))
    r_sig2 = p_sign.add_run("Erode District Collectorate • Government of Tamil Nadu")
    _format_run_fonts(r_sig2, size_pt=8, color_rgb=(71, 85, 105))

    p_bot = doc.add_paragraph()
    p_bot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bot.paragraph_format.space_before = Pt(8)
    r_bot = p_bot.add_run("Official Certified Record • AI Administrative Co-Pilot (GDP Assistant) • Erode District")
    _format_run_fonts(r_bot, size_pt=7.5, color_rgb=(148, 163, 184))


# ==============================================================================
# REPORTLAB (PDF) NUMBERED CANVAS
# ==============================================================================

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        # Top banner on first page
        if self._pageNumber == 1:
            self.setFillColor(RL_NAVY)
            self.rect(0, 842 - 50, 595, 50, fill=1, stroke=0)
            self.setFillColor(RL_GOLD)
            self.rect(0, 842 - 53, 595, 3, fill=1, stroke=0)

            self.setFillColor(colors.white)
            self.setFont(PDF_BOLD_FONT, 10.5)
            self.drawCentredString(297.5, 842 - 19, "GOVERNMENT OF TAMIL NADU")
            self.setFont(PDF_REGULAR_FONT, 7.5)
            self.drawCentredString(297.5, 842 - 31, "REVENUE & DISASTER MANAGEMENT DEPARTMENT • ERODE DISTRICT COLLECTORATE")
            self.setFont(PDF_BOLD_FONT, 8)
            self.setFillColor(colors.HexColor("#EADBC8"))
            self.drawCentredString(297.5, 842 - 43, "AI ADMINISTRATIVE CO-PILOT (GDP ASSISTANT) — CERTIFIED SYSTEM AUDIT REPORT")

        # Footer on all pages
        self.setStrokeColor(RL_BORDER)
        self.setLineWidth(0.5)
        self.line(36, 30, 559, 30)

        self.setFont(PDF_REGULAR_FONT, 7)
        self.setFillColor(RL_SLATE)
        self.drawString(36, 18, "Official Government Document • District Collectorate Erode • Public Grievance Pre-Processing Cell")
        self.drawRightString(559, 18, f"Page {self._pageNumber} of {page_count}")
        self.restoreState()


# ==============================================================================
# 1. OFFICER PERFORMANCE (PDF & DOCX)
# ==============================================================================

def generate_pdf_officer_performance(data: Dict[str, Any], selected_officer_id: str = "all") -> bytes:
    meta = data.get("reportMetadata", {})
    officers = data.get("officersDirectory", [])
    petitions = data.get("petitionProcessingHistory", [])

    if selected_officer_id and selected_officer_id != "all":
        officers = [o for o in officers if str(o.get("id")) == str(selected_officer_id)]

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=36, rightMargin=36, topMargin=60, bottomMargin=45
    )
    styles = _get_pdf_styles()
    elements = []

    # Metadata Block Table
    meta_table_data = [
        [
            Paragraph(f"<b>Report Title:</b> Officer Processed Petition Counts & Status", styles['normal']),
            Paragraph(f"<b>Generated By:</b> {meta.get('generatedBy', 'District Administrator')}", styles['normal'])
        ],
        [
            Paragraph("<b>District / Office:</b> Erode District Collectorate", styles['normal']),
            Paragraph("<b>Security Level:</b> <font color='#166534'><b>OFFICIAL / RESTRICTED</b></font>", styles['normal'])
        ],
        [
            Paragraph(f"<b>Generated At:</b> {meta.get('generatedAt', '')[:19].replace('T', ' ')}", styles['normal']),
            Paragraph(f"<b>Officer Scope:</b> {selected_officer_id if selected_officer_id != 'all' else 'All Registered Officers'}", styles['normal'])
        ]
    ]
    meta_table = Table(meta_table_data, colWidths=[260, 260])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), RL_LIGHT_BG),
        ('BOX', (0, 0), (-1, -1), 1, RL_BORDER),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    elements.append(meta_table)
    elements.append(Spacer(1, 10))

    # KPI Summary Cards
    kpi_data = [
        [
            Paragraph(f"<font size=12 color='#102C57'><b>{meta.get('totalOfficers', len(officers))}</b></font><br/><font size=7 color='#475569'>Total Officers</font>", styles['normal']),
            Paragraph(f"<font size=12 color='#102C57'><b>{meta.get('activeOfficers', 0)}</b></font><br/><font size=7 color='#475569'>Active Sessions</font>", styles['normal']),
            Paragraph(f"<font size=12 color='#102C57'><b>{meta.get('totalPetitions', len(petitions))}</b></font><br/><font size=7 color='#475569'>Total Petitions</font>", styles['normal']),
            Paragraph(f"<font size=12 color='#166534'><b>{meta.get('totalApproved', 0)}</b></font><br/><font size=7 color='#475569'>Approved / Resolved</font>", styles['normal'])
        ]
    ]
    kpi_table = Table(kpi_data, colWidths=[130, 130, 130, 130])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.white),
        ('BOX', (0, 0), (-1, -1), 0.8, RL_BORDER),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, RL_BORDER),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(kpi_table)
    elements.append(Spacer(1, 12))

    # Section 1: Officers Roster & Performance Table
    elements.append(Paragraph("<b>1. Officer Directory & Petition Resolution Metrics</b>", styles['heading']))
    elements.append(Spacer(1, 4))

    perf_head = [
        Paragraph("<b>Officer ID</b>", styles['bold']),
        Paragraph("<b>Official Name</b>", styles['bold']),
        Paragraph("<b>Department</b>", styles['bold']),
        Paragraph("<b>Total</b>", styles['bold']),
        Paragraph("<b>Approved</b>", styles['bold']),
        Paragraph("<b>In Review</b>", styles['bold']),
        Paragraph("<b>Pending</b>", styles['bold']),
        Paragraph("<b>Success %</b>", styles['bold'])
    ]
    perf_rows = [perf_head]
    for o in officers:
        tamil_name = f"<br/><font size=6.5 color='#475569'>({o.get('nameTamil', '')})</font>" if o.get('nameTamil') else ""
        perf_rows.append([
            Paragraph(str(o.get("id", "—")), styles['bold']),
            Paragraph(f"<b>{o.get('name', '—')}</b>{tamil_name}", styles['tamil']),
            Paragraph(f"{o.get('department', '—')}<br/><font size=6.5 color='#64748B'>{o.get('designation', '—')}</font>", styles['normal']),
            Paragraph(str(o.get("totalProcessed", 0)), styles['bold']),
            Paragraph(f"<font color='#166534'><b>{o.get('approved', 0)}</b></font>", styles['normal']),
            Paragraph(str(o.get("inProgress", 0)), styles['normal']),
            Paragraph(str(o.get("pending", 0)), styles['normal']),
            Paragraph(f"<font color='#166534'><b>{o.get('successRate', '100%')}</b></font>", styles['bold'])
        ])

    perf_table = Table(perf_rows, colWidths=[65, 110, 125, 42, 45, 45, 40, 48])
    perf_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), RL_NAVY),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, RL_LIGHT_BG]),
        ('GRID', (0, 0), (-1, -1), 0.5, RL_BORDER),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(perf_table)

    doc.build(elements, canvasmaker=NumberedCanvas)
    return buffer.getvalue()


def generate_docx_officer_performance(data: Dict[str, Any], selected_officer_id: str = "all") -> bytes:
    meta = data.get("reportMetadata", {})
    officers = data.get("officersDirectory", [])

    if selected_officer_id and selected_officer_id != "all":
        officers = [o for o in officers if str(o.get("id")) == str(selected_officer_id)]

    doc = Document()
    _add_docx_header(doc, "OFFICER PETITION PROCESSING & RESOLUTION AUDIT REPORT")

    p1 = doc.add_paragraph()
    r_h1 = p1.add_run("1. Executive Summary & Officer Performance Roster")
    _format_run_fonts(r_h1, size_pt=10.5, bold=True, color_rgb=(16, 44, 87))

    table = doc.add_table(rows=1, cols=8)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(table)

    headers = ["Officer ID", "Official Name", "Department / Role", "Total", "Approved", "In Review", "Pending", "Resolution %"]
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        r = p.add_run(h)
        _format_run_fonts(r, size_pt=8.5, bold=True, color_rgb=(255, 255, 255))
        _set_cell_bg(hdr_cells[i], HEX_NAVY)
        _set_cell_margins(hdr_cells[i], top=100, bottom=100, left=100, right=100)

    for idx, off in enumerate(officers):
        row_cells = table.add_row().cells
        bg = HEX_LIGHT_BG if idx % 2 == 1 else "FFFFFF"
        for c in row_cells:
            _set_cell_bg(c, bg)
            _set_cell_margins(c, top=80, bottom=80, left=100, right=100)

        # ID
        r0 = row_cells[0].paragraphs[0].add_run(str(off.get("id", "—")))
        _format_run_fonts(r0, size_pt=8, bold=True, color_rgb=(16, 44, 87))

        # Name + Tamil
        p_name = row_cells[1].paragraphs[0]
        r_n1 = p_name.add_run(f"{off.get('name', '—')}\n")
        _format_run_fonts(r_n1, size_pt=8.5, bold=True)
        if off.get('nameTamil'):
            r_n2 = p_name.add_run(f"({off.get('nameTamil')})")
            _format_run_fonts(r_n2, size_pt=7.5, color_rgb=(71, 85, 105))

        # Dept
        p_dept = row_cells[2].paragraphs[0]
        r_d1 = p_dept.add_run(f"{off.get('department', '—')}\n")
        _format_run_fonts(r_d1, size_pt=8)
        r_d2 = p_dept.add_run(str(off.get('designation', '—')))
        _format_run_fonts(r_d2, size_pt=7, color_rgb=(100, 116, 139))

        # Counts
        for c_idx, val, is_b, col in [
            (3, str(off.get("totalProcessed", 0)), True, (16, 44, 87)),
            (4, str(off.get("approved", 0)), True, (22, 101, 52)),
            (5, str(off.get("inProgress", 0)), False, (15, 23, 42)),
            (6, str(off.get("pending", 0)), False, (180, 83, 9)),
            (7, str(off.get("successRate", "100%")), True, (22, 101, 52)),
        ]:
            r_val = row_cells[c_idx].paragraphs[0].add_run(val)
            _format_run_fonts(r_val, size_pt=8, bold=is_b, color_rgb=col)

    _add_docx_footer_and_seal(doc)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ==============================================================================
# 2. AUDIT LOGS (PDF & DOCX)
# ==============================================================================

def generate_pdf_audit_logs(data: Dict[str, Any], filter_officer_id: str = "all") -> bytes:
    logs = data.get("recentAuditLogs", [])
    if filter_officer_id and filter_officer_id != "all":
        logs = [l for l in logs if str(l.get("officer_id", "")).lower() == str(filter_officer_id).lower()]

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=36, rightMargin=36, topMargin=60, bottomMargin=45
    )
    styles = _get_pdf_styles()
    elements = []

    scope_str = "All Officers & Modules" if filter_officer_id == "all" else f"Officer ID: {filter_officer_id}"
    elements.append(Paragraph(f"<b>Chronological System & Officer Audit History Trail</b> ({scope_str})", styles['heading']))
    elements.append(Spacer(1, 6))

    head = [
        Paragraph("<b>Timestamp</b>", styles['bold']),
        Paragraph("<b>Officer ID</b>", styles['bold']),
        Paragraph("<b>Category</b>", styles['bold']),
        Paragraph("<b>Action Type</b>", styles['bold']),
        Paragraph("<b>Audit Event Details</b>", styles['bold'])
    ]
    rows = [head]
    for l in logs[:150]:
        ts = str(l.get("timestamp", ""))[:19].replace("T", " ")
        rows.append([
            Paragraph(ts, styles['normal']),
            Paragraph(str(l.get("officer_id", "SYSTEM")), styles['bold']),
            Paragraph(str(l.get("category", "GDP Assistant")), styles['normal']),
            Paragraph(str(l.get("action", "EVENT")), styles['normal']),
            Paragraph(str(l.get("details", "—")), styles['tamil'])
        ])

    table = Table(rows, colWidths=[90, 75, 85, 75, 195])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), RL_NAVY),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, RL_LIGHT_BG]),
        ('GRID', (0, 0), (-1, -1), 0.5, RL_BORDER),
        ('TOPPADDING', (0, 0), (-1, -1), 3.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3.5),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(table)

    doc.build(elements, canvasmaker=NumberedCanvas)
    return buffer.getvalue()


def generate_docx_audit_logs(data: Dict[str, Any], filter_officer_id: str = "all") -> bytes:
    logs = data.get("recentAuditLogs", [])
    if filter_officer_id and filter_officer_id != "all":
        logs = [l for l in logs if str(l.get("officer_id", "")).lower() == str(filter_officer_id).lower()]

    doc = Document()
    _add_docx_header(doc, "SYSTEM & OFFICER AUDIT HISTORY TRAIL REPORT")

    p1 = doc.add_paragraph()
    scope_str = "All Officers" if filter_officer_id == "all" else f"Officer: {filter_officer_id}"
    r_h1 = p1.add_run(f"1. Chronological Audit Log History ({scope_str} • {len(logs)} Events)")
    _format_run_fonts(r_h1, size_pt=10.5, bold=True, color_rgb=(16, 44, 87))

    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(table)

    headers = ["Timestamp", "Officer ID", "Category", "Action Type", "Audit Event Details"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        _format_run_fonts(r, size_pt=8.5, bold=True, color_rgb=(255, 255, 255))
        _set_cell_bg(cell, HEX_NAVY)
        _set_cell_margins(cell, top=100, bottom=100, left=100, right=100)

    for idx, l in enumerate(logs[:180]):
        row = table.add_row().cells
        bg = HEX_LIGHT_BG if idx % 2 == 1 else "FFFFFF"
        for c in row:
            _set_cell_bg(c, bg)
            _set_cell_margins(c, top=70, bottom=70, left=90, right=90)

        r0 = row[0].paragraphs[0].add_run(str(l.get("timestamp", ""))[:19].replace("T", " "))
        _format_run_fonts(r0, size_pt=7.5)

        r1 = row[1].paragraphs[0].add_run(str(l.get("officer_id", "SYSTEM")))
        _format_run_fonts(r1, size_pt=8, bold=True, color_rgb=(16, 44, 87))

        r2 = row[2].paragraphs[0].add_run(str(l.get("category", "GDP Assistant")))
        _format_run_fonts(r2, size_pt=8)

        r3 = row[3].paragraphs[0].add_run(str(l.get("action", "EVENT")))
        _format_run_fonts(r3, size_pt=8)

        r4 = row[4].paragraphs[0].add_run(str(l.get("details", "—")))
        _format_run_fonts(r4, size_pt=8)

    _add_docx_footer_and_seal(doc)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ==============================================================================
# 3. SINGLE PETITION DOSSIER (PDF & DOCX)
# ==============================================================================

def generate_pdf_single_petition(petition: Dict[str, Any]) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=36, rightMargin=36, topMargin=60, bottomMargin=45
    )
    styles = _get_pdf_styles()
    elements = []

    pet_num = petition.get("petitionNumber") or petition.get("id") or "PET-RECORD"

    # Header Card
    p_hdr = [
        [
            Paragraph(f"<font size=11 color='#102C57'><b>Petition Dossier: {pet_num}</b></font><br/>"
                      f"<font size=8 color='#475569'>Status: <b>{petition.get('status', 'Approved')}</b> | Priority: {petition.get('priority', 'MEDIUM')}</font>", styles['normal']),
            Paragraph(f"<font size=8><b>Assigned Officer:</b> {petition.get('officerName', 'Assigned Officer')}<br/>"
                      f"<b>Submission Date:</b> {str(petition.get('createdAt', ''))[:10]}</font>", styles['normal'])
        ]
    ]
    hdr_table = Table(p_hdr, colWidths=[260, 260])
    hdr_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), RL_LIGHT_BG),
        ('BOX', (0, 0), (-1, -1), 1, RL_NAVY),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    elements.append(hdr_table)
    elements.append(Spacer(1, 10))

    # 1. Applicant Details
    elements.append(Paragraph("<b>1. Applicant Details & Location Hierarchy</b>", styles['heading']))
    elements.append(Spacer(1, 3))

    app_data = [
        [
            Paragraph("<b>Petitioner Name</b>", styles['bold']),
            Paragraph(str(petition.get("applicantName", "—")), styles['tamil_bold']),
            Paragraph("<b>District</b>", styles['bold']),
            Paragraph(str(petition.get("district", "Erode")), styles['normal'])
        ],
        [
            Paragraph("<b>Mobile Number</b>", styles['bold']),
            Paragraph(str(petition.get("mobile", "—")), styles['normal']),
            Paragraph("<b>Taluk / Firka</b>", styles['bold']),
            Paragraph(f"{petition.get('taluk', 'Erode')} / {petition.get('firka', 'Erode Urban')}", styles['tamil'])
        ],
        [
            Paragraph("<b>Email Address</b>", styles['bold']),
            Paragraph(str(petition.get("email", "—")), styles['normal']),
            Paragraph("<b>Village / Block</b>", styles['bold']),
            Paragraph(f"{petition.get('village', 'Surampatti')} / {petition.get('block', 'Erode')}", styles['tamil'])
        ],
        [
            Paragraph("<b>Residential Address</b>", styles['bold']),
            Paragraph(str(petition.get("address", "Erode, Tamil Nadu")), styles['tamil']),
            Paragraph("<b>Intake Channel</b>", styles['bold']),
            Paragraph(str(petition.get("intakeChannel", "Collectorate Counter")), styles['normal'])
        ]
    ]
    t1 = Table(app_data, colWidths=[110, 150, 110, 150])
    t1.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), RL_LABEL_BG),
        ('BACKGROUND', (2, 0), (2, -1), RL_LABEL_BG),
        ('GRID', (0, 0), (-1, -1), 0.5, RL_BORDER),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(t1)
    elements.append(Spacer(1, 10))

    # 2. Classification
    elements.append(Paragraph("<b>2. Administrative Grievance Classification</b>", styles['heading']))
    elements.append(Spacer(1, 3))

    class_data = [
        [
            Paragraph("<b>Department</b>", styles['bold']),
            Paragraph(str(petition.get("department", "Revenue Administration")), styles['normal']),
            Paragraph("<b>Grievance Category</b>", styles['bold']),
            Paragraph(str(petition.get("category", "Patta & Land Records")), styles['tamil'])
        ],
        [
            Paragraph("<b>Sub-Category</b>", styles['bold']),
            Paragraph(str(petition.get("subCategory", "Patta Transfer")), styles['tamil']),
            Paragraph("<b>Officer ID</b>", styles['bold']),
            Paragraph(str(petition.get("officerId", "ADM-ERODE-001")), styles['bold'])
        ]
    ]
    t2 = Table(class_data, colWidths=[110, 150, 110, 150])
    t2.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), RL_LABEL_BG),
        ('BACKGROUND', (2, 0), (2, -1), RL_LABEL_BG),
        ('GRID', (0, 0), (-1, -1), 0.5, RL_BORDER),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(t2)
    elements.append(Spacer(1, 10))

    # 3. AI Summary & Action Items
    elements.append(Paragraph("<b>3. AI Grievance Synopsis & Action Items</b>", styles['heading']))
    elements.append(Spacer(1, 3))

    action_items_raw = petition.get("actionItems", [])
    formatted_actions = []
    if isinstance(action_items_raw, list):
        for act in action_items_raw:
            if isinstance(act, dict):
                act_text = act.get("action") or act.get("description") or act.get("task") or json.dumps(act)
                formatted_actions.append(str(act_text))
            else:
                formatted_actions.append(str(act))
    elif isinstance(action_items_raw, str):
        formatted_actions.append(action_items_raw)

    if not formatted_actions:
        formatted_actions = [
            "வருவாய் ஆய்வாளர் மூலம் கள ஆய்வு மேற்கொள்ளுதல்",
            "கிராம நிர்வாக ஆவணங்கள் மற்றும் பதிவேடுகளை சரிபார்த்தல்",
            "மாவட்ட வருவாய் அலுவலர் (DRO) இறுதி ஆணை பிறப்பித்தல்"
        ]

    action_items_str = "<br/>• " + "<br/>• ".join(formatted_actions)

    summary_data = [
        [
            Paragraph("<b>Tamil Summary<br/>(மனு சுருக்கம்)</b>", styles['bold']),
            Paragraph(str(petition.get("summaryTamil") or petition.get("description", "மனு விவரங்கள் பதிவு செய்யப்பட்டுள்ளன.")), styles['tamil'])
        ],
        [
            Paragraph("<b>English Synopsis</b>", styles['bold']),
            Paragraph(str(petition.get("summaryEnglish") or "Grievance received and recorded in public registry for administrative action."), styles['normal'])
        ],
        [
            Paragraph("<b>Action Points</b>", styles['bold']),
            Paragraph(action_items_str, styles['tamil'])
        ]
    ]
    t3 = Table(summary_data, colWidths=[130, 390])
    t3.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), RL_LABEL_BG),
        ('GRID', (0, 0), (-1, -1), 0.5, RL_BORDER),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(t3)

    doc.build(elements, canvasmaker=NumberedCanvas)
    return buffer.getvalue()


def generate_docx_single_petition(petition: Dict[str, Any]) -> bytes:
    doc = Document()
    pet_num = petition.get("petitionNumber") or petition.get("id") or "PET-RECORD"
    _add_docx_header(doc, f"COMPREHENSIVE PETITION DOSSIER — {pet_num}")

    # Top Status Pill Strip
    p_status = doc.add_paragraph()
    r_stat1 = p_status.add_run(f"Petition Number: {pet_num}  |  ")
    _format_run_fonts(r_stat1, size_pt=9.5, bold=True, color_rgb=(16, 44, 87))
    r_stat2 = p_status.add_run(f"Status: {petition.get('status', 'Approved')}  |  Priority: {petition.get('priority', 'MEDIUM')}")
    _format_run_fonts(r_stat2, size_pt=9, bold=True, color_rgb=(22, 101, 52))

    # 1. Applicant Identity Table
    p1 = doc.add_paragraph()
    r_h1 = p1.add_run("1. Applicant Details & Geographic Hierarchy")
    _format_run_fonts(r_h1, size_pt=10, bold=True, color_rgb=(16, 44, 87))

    t1 = doc.add_table(rows=4, cols=4)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(t1)

    app_fields = [
        ("Petitioner Name:", str(petition.get("applicantName", "—")), "District:", str(petition.get("district", "Erode"))),
        ("Mobile Number:", str(petition.get("mobile", "—")), "Taluk / Firka:", f"{petition.get('taluk', 'Erode')} / {petition.get('firka', 'Erode Urban')}"),
        ("Email Address:", str(petition.get("email", "—")), "Village / Block:", f"{petition.get('village', 'Surampatti')} / {petition.get('block', 'Erode')}"),
        ("Residential Address:", str(petition.get("address", "Erode District, Tamil Nadu")), "Intake Channel:", str(petition.get("intakeChannel", "Collectorate Counter"))),
    ]

    for row_idx, (l1, v1, l2, v2) in enumerate(app_fields):
        row = t1.rows[row_idx].cells
        for c in row:
            _set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        
        _set_cell_bg(row[0], HEX_LABEL_BG)
        _set_cell_bg(row[2], HEX_LABEL_BG)

        r0 = row[0].paragraphs[0].add_run(l1)
        _format_run_fonts(r0, size_pt=8.5, bold=True, color_rgb=(71, 85, 105))

        r1 = row[1].paragraphs[0].add_run(v1)
        _format_run_fonts(r1, size_pt=8.5, bold=(row_idx == 0))

        r2 = row[2].paragraphs[0].add_run(l2)
        _format_run_fonts(r2, size_pt=8.5, bold=True, color_rgb=(71, 85, 105))

        r3 = row[3].paragraphs[0].add_run(v2)
        _format_run_fonts(r3, size_pt=8.5)

    # 2. Administrative Classification Table
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(8)
    r_h2 = p2.add_run("2. Administrative Grievance Classification")
    _format_run_fonts(r_h2, size_pt=10, bold=True, color_rgb=(16, 44, 87))

    t2 = doc.add_table(rows=2, cols=4)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(t2)

    class_fields = [
        ("Department:", str(petition.get("department", "Revenue Administration")), "Grievance Category:", str(petition.get("category", "Patta & Land Records"))),
        ("Sub-Category:", str(petition.get("subCategory", "Patta Transfer")), "Assigned Officer:", f"{petition.get('officerName', 'Assigned Officer')} ({petition.get('officerId', 'ADM-ERODE-001')})"),
    ]

    for row_idx, (l1, v1, l2, v2) in enumerate(class_fields):
        row = t2.rows[row_idx].cells
        for c in row:
            _set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        
        _set_cell_bg(row[0], HEX_LABEL_BG)
        _set_cell_bg(row[2], HEX_LABEL_BG)

        r0 = row[0].paragraphs[0].add_run(l1)
        _format_run_fonts(r0, size_pt=8.5, bold=True, color_rgb=(71, 85, 105))

        r1 = row[1].paragraphs[0].add_run(v1)
        _format_run_fonts(r1, size_pt=8.5, bold=(row_idx == 0))

        r2 = row[2].paragraphs[0].add_run(l2)
        _format_run_fonts(r2, size_pt=8.5, bold=True, color_rgb=(71, 85, 105))

        r3 = row[3].paragraphs[0].add_run(v2)
        _format_run_fonts(r3, size_pt=8.5)

    # 3. AI Grievance Synopsis & Action Points
    p3 = doc.add_paragraph()
    p3.paragraph_format.space_before = Pt(8)
    r_h3 = p3.add_run("3. AI Grievance Synopsis & Action Items")
    _format_run_fonts(r_h3, size_pt=10, bold=True, color_rgb=(16, 44, 87))

    t3 = doc.add_table(rows=3, cols=2)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(t3)

    action_items_raw = petition.get("actionItems", [])
    formatted_actions = []
    if isinstance(action_items_raw, list):
        for act in action_items_raw:
            if isinstance(act, dict):
                act_text = act.get("action") or act.get("description") or act.get("task") or json.dumps(act)
                formatted_actions.append(str(act_text))
            else:
                formatted_actions.append(str(act))
    elif isinstance(action_items_raw, str):
        formatted_actions.append(action_items_raw)

    if not formatted_actions:
        formatted_actions = [
            "வருவாய் ஆய்வாளர் மூலம் கள ஆய்வு மேற்கொள்ளுதல்",
            "கிராம நிர்வாக ஆவணங்கள் மற்றும் பட்டா பதிவேடுகளை சரிபார்த்தல்",
            "மாவட்ட வருவாய் அலுவலர் (DRO) இறுதி ஆணை பிறப்பித்தல்"
        ]

    synopsis_rows = [
        ("Tamil Summary\n(மனு சுருக்கம்):", str(petition.get("summaryTamil") or petition.get("description", "மனு விவரங்கள் பதிவு செய்யப்பட்டுள்ளன."))),
        ("English Synopsis:", str(petition.get("summaryEnglish") or "Grievance recorded into the public registry for administrative action.")),
        ("Action Items:", "\n".join([f"• {a}" for a in formatted_actions]))
    ]

    for row_idx, (label, content) in enumerate(synopsis_rows):
        row = t3.rows[row_idx].cells
        for c in row:
            _set_cell_margins(c, top=90, bottom=90, left=110, right=110)
        _set_cell_bg(row[0], HEX_LABEL_BG)

        r_lbl = row[0].paragraphs[0].add_run(label)
        _format_run_fonts(r_lbl, size_pt=8.5, bold=True, color_rgb=(71, 85, 105))

        r_val = row[1].paragraphs[0].add_run(content)
        _format_run_fonts(r_val, size_pt=8.5)

    _add_docx_footer_and_seal(doc)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ==============================================================================
# 4. TOTAL PETITIONS RECEIVED & INTAKE ANALYSIS (PDF & DOCX)
# ==============================================================================

def generate_pdf_intake_report(data: Dict[str, Any]) -> bytes:
    meta = data.get("reportMetadata", {})
    analysis = data.get("intakeAnalysis", {})
    by_channel = analysis.get("byChannel", {})
    by_dept = analysis.get("byDepartment", {})

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=36, rightMargin=36, topMargin=60, bottomMargin=45
    )
    styles = _get_pdf_styles()
    elements = []

    elements.append(Paragraph("<b>Total Petitions Received & Multi-Channel Intake Report</b>", styles['heading']))
    elements.append(Paragraph(f"<font size=8 color='#475569'>Total Petitions Ingested: <b>{meta.get('totalPetitions', 0)}</b> | Approved: <b>{meta.get('totalApproved', 0)}</b> | Pending: <b>{meta.get('totalPending', 0)}</b></font>", styles['normal']))
    elements.append(Spacer(1, 8))

    # Channels Table
    elements.append(Paragraph("<b>1. CM Grievance Ingestion Channels Breakdown</b>", styles['heading']))
    elements.append(Spacer(1, 3))
    ch_rows = [[Paragraph("<b>Intake Channel Name</b>", styles['bold']), Paragraph("<b>Petitions Count</b>", styles['bold'])]]
    for ch, cnt in by_channel.items():
        ch_rows.append([Paragraph(str(ch), styles['normal']), Paragraph(str(cnt), styles['bold'])])
    t_ch = Table(ch_rows, colWidths=[380, 140])
    t_ch.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), RL_NAVY),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, RL_LIGHT_BG]),
        ('GRID', (0, 0), (-1, -1), 0.5, RL_BORDER),
        ('TOPPADDING', (0, 0), (-1, -1), 3.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3.5),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(t_ch)
    elements.append(Spacer(1, 8))

    # Departments Table
    elements.append(Paragraph("<b>2. Distribution Across Government Departments</b>", styles['heading']))
    elements.append(Spacer(1, 3))
    dept_rows = [[Paragraph("<b>Department Name</b>", styles['bold']), Paragraph("<b>Total Grievances</b>", styles['bold'])]]
    for dp, cnt in by_dept.items():
        dept_rows.append([Paragraph(str(dp), styles['normal']), Paragraph(str(cnt), styles['bold'])])
    t_dp = Table(dept_rows, colWidths=[380, 140])
    t_dp.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), RL_NAVY),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, RL_LIGHT_BG]),
        ('GRID', (0, 0), (-1, -1), 0.5, RL_BORDER),
        ('TOPPADDING', (0, 0), (-1, -1), 3.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3.5),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(t_dp)

    doc.build(elements, canvasmaker=NumberedCanvas)
    return buffer.getvalue()


def generate_docx_intake_report(data: Dict[str, Any]) -> bytes:
    meta = data.get("reportMetadata", {})
    analysis = data.get("intakeAnalysis", {})
    by_channel = analysis.get("byChannel", {})
    by_dept = analysis.get("byDepartment", {})

    doc = Document()
    _add_docx_header(doc, "TOTAL PETITIONS RECEIVED & INTAKE ANALYSIS")

    p1 = doc.add_paragraph()
    r_h1 = p1.add_run(f"1. Ingestion Breakdown Across 21 CM Grievance Channels (Total: {meta.get('totalPetitions', 0)})")
    _format_run_fonts(r_h1, size_pt=10, bold=True, color_rgb=(16, 44, 87))

    t_ch = doc.add_table(rows=1, cols=2)
    t_ch.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(t_ch)

    hdr_c = t_ch.rows[0].cells
    hdr_c[0].text = ""
    hdr_c[1].text = ""
    r0 = hdr_c[0].paragraphs[0].add_run("CM Grievance Intake Channel")
    _format_run_fonts(r0, size_pt=8.5, bold=True, color_rgb=(255, 255, 255))
    r1 = hdr_c[1].paragraphs[0].add_run("Petitions Received")
    _format_run_fonts(r1, size_pt=8.5, bold=True, color_rgb=(255, 255, 255))
    _set_cell_bg(hdr_c[0], HEX_NAVY)
    _set_cell_bg(hdr_c[1], HEX_NAVY)
    _set_cell_margins(hdr_c[0], top=100, bottom=100, left=120, right=120)
    _set_cell_margins(hdr_c[1], top=100, bottom=100, left=120, right=120)

    for idx, (ch, cnt) in enumerate(by_channel.items()):
        row = t_ch.add_row().cells
        bg = HEX_LIGHT_BG if idx % 2 == 1 else "FFFFFF"
        for c in row:
            _set_cell_bg(c, bg)
            _set_cell_margins(c, top=70, bottom=70, left=100, right=100)

        r_c0 = row[0].paragraphs[0].add_run(ch)
        _format_run_fonts(r_c0, size_pt=8.5)

        r_c1 = row[1].paragraphs[0].add_run(str(cnt))
        _format_run_fonts(r_c1, size_pt=8.5, bold=True, color_rgb=(16, 44, 87))

    # Departments Table
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(12)
    r_h2 = p2.add_run("2. Distribution Across Government Departments")
    _format_run_fonts(r_h2, size_pt=10, bold=True, color_rgb=(16, 44, 87))

    t_dp = doc.add_table(rows=1, cols=2)
    t_dp.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(t_dp)

    hdr_d = t_dp.rows[0].cells
    hdr_d[0].text = ""
    hdr_d[1].text = ""
    r_d0 = hdr_d[0].paragraphs[0].add_run("Department Name")
    _format_run_fonts(r_d0, size_pt=8.5, bold=True, color_rgb=(255, 255, 255))
    r_d1 = hdr_d[1].paragraphs[0].add_run("Total Grievances")
    _format_run_fonts(r_d1, size_pt=8.5, bold=True, color_rgb=(255, 255, 255))
    _set_cell_bg(hdr_d[0], HEX_NAVY)
    _set_cell_bg(hdr_d[1], HEX_NAVY)
    _set_cell_margins(hdr_d[0], top=100, bottom=100, left=120, right=120)
    _set_cell_margins(hdr_d[1], top=100, bottom=100, left=120, right=120)

    for idx, (dp, cnt) in enumerate(by_dept.items()):
        row = t_dp.add_row().cells
        bg = HEX_LIGHT_BG if idx % 2 == 1 else "FFFFFF"
        for c in row:
            _set_cell_bg(c, bg)
            _set_cell_margins(c, top=70, bottom=70, left=100, right=100)

        r_dp0 = row[0].paragraphs[0].add_run(dp)
        _format_run_fonts(r_dp0, size_pt=8.5)

        r_dp1 = row[1].paragraphs[0].add_run(str(cnt))
        _format_run_fonts(r_dp1, size_pt=8.5, bold=True, color_rgb=(16, 44, 87))

    _add_docx_footer_and_seal(doc)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ==============================================================================
# UNIFIED DISPATCHER
# ==============================================================================

def generate_report_document(
    template: str,
    fmt: str,
    data: Dict[str, Any],
    officer_id: str = "all",
    petition_id: Optional[str] = None
) -> Tuple[bytes, str, str]:
    """
    Main dispatcher for creating certified report files.
    Returns: (file_bytes, filename, media_type)
    """
    fmt = fmt.lower().strip()
    date_str = datetime.now(timezone.utc).strftime("%Y%m%d")

    if template == "officer_performance":
        if fmt == "docx":
            content = generate_docx_officer_performance(data, officer_id)
            return content, f"Officer_Performance_Report_{date_str}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        else:
            content = generate_pdf_officer_performance(data, officer_id)
            return content, f"Officer_Performance_Report_{date_str}.pdf", "application/pdf"

    elif template == "audit_logs":
        scope = "All" if officer_id == "all" else str(officer_id)
        if fmt == "docx":
            content = generate_docx_audit_logs(data, officer_id)
            return content, f"System_Audit_Log_{scope}_{date_str}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        else:
            content = generate_pdf_audit_logs(data, officer_id)
            return content, f"System_Audit_Log_{scope}_{date_str}.pdf", "application/pdf"

    elif template == "single_petition":
        petitions = data.get("petitionProcessingHistory", [])
        target = next((p for p in petitions if str(p.get("id")) == str(petition_id) or str(p.get("petitionNumber")) == str(petition_id)), None)
        if not target and petitions:
            target = petitions[0]
        if not target:
            target = {"petitionNumber": petition_id or "PET-001", "applicantName": "Grievance Applicant", "status": "Approved"}

        p_num = str(target.get("petitionNumber") or "PET")
        if fmt == "docx":
            content = generate_docx_single_petition(target)
            return content, f"Petition_Dossier_{p_num}_{date_str}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        else:
            content = generate_pdf_single_petition(target)
            return content, f"Petition_Dossier_{p_num}_{date_str}.pdf", "application/pdf"

    elif template == "intake_report":
        if fmt == "docx":
            content = generate_docx_intake_report(data)
            return content, f"Total_Petitions_Received_Report_{date_str}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        else:
            content = generate_pdf_intake_report(data)
            return content, f"Total_Petitions_Received_Report_{date_str}.pdf", "application/pdf"

    else:
        # Default to officer performance
        content = generate_pdf_officer_performance(data, officer_id)
        return content, f"Audit_Report_{date_str}.pdf", "application/pdf"
